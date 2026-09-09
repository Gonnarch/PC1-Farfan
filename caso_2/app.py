from pathlib import Path
from io import BytesIO
import os
import re
import sqlite3
from flask import Flask, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 16384
DATA_DIR = Path(os.environ.get('DATA_DIR', str(Path(__file__).parent / 'data'))).resolve()
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB = DATA_DIR / 'consultas.sqlite3'
OFFICIAL_URL = 'https://consultaelectoral.onpe.gob.pe/inicio'
# Conserva la tabla existente para no borrar registros del usuario.
with sqlite3.connect(DB) as con:
    con.execute('CREATE TABLE IF NOT EXISTS consultas (dni TEXT PRIMARY KEY, estado TEXT NOT NULL DEFAULT "Pendiente", region TEXT DEFAULT "", provincia TEXT DEFAULT "", distrito TEXT DEFAULT "", direccion TEXT DEFAULT "", fecha TEXT DEFAULT "")')

def normalizar_dni(value):
    if isinstance(value, int):
        value = str(value).zfill(8)
    value = str(value or '').strip()
    if not re.fullmatch('[0-9]{8}', value):
        raise ValueError('El DNI debe tener 8 dígitos.')
    return value

def rows():
    with sqlite3.connect(DB) as con:
        con.row_factory = sqlite3.Row
        return [dict(row) for row in con.execute('SELECT * FROM consultas ORDER BY dni')]

@app.get('/health')
def health():
    return {'status': 'ok'}

def excel(headers, data, filename):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Miembros de mesa' if len(headers) > 1 else 'DNIs'
    sheet.append(headers)
    for row in data:
        sheet.append(row)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            if cell.column == 1:
                cell.number_format = '@'
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='17365D')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 44 if col[0].value == 'Dirección del local de votación' else 24
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=filename, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

def page(mensaje='', error=False, status=200):
    miembros = [r for r in rows() if r['estado'] == 'Sí']
    return render_template('index.html', title='Miembros de mesa', subtitle='Caso 2 / Miembros de mesa', mensaje=mensaje, error=error, rows=miembros, official_url=OFFICIAL_URL), status

@app.get('/')
def inicio():
    return page()

@app.post('/guardar')
def guardar():
    try:
        dni = normalizar_dni(request.form.get('dni'))
        fields = [request.form.get(k, '').strip() for k in ('region','provincia','distrito','direccion')]
        if not all(fields):
            raise ValueError('Completa región, provincia, distrito y dirección del local de votación.')
        if any(len(value) > 300 for value in fields):
            raise ValueError('Cada campo admite como máximo 300 caracteres.')
        if request.form.get('miembro') != 'si':
            raise ValueError('Registra solamente a quienes figuran como miembros de mesa en el portal oficial.')
        with sqlite3.connect(DB) as con:
            con.execute('INSERT INTO consultas (dni, estado, region, provincia, distrito, direccion, fecha) VALUES (?, ?, ?, ?, ?, ?, ?) ON CONFLICT(dni) DO UPDATE SET estado=excluded.estado, region=excluded.region, provincia=excluded.provincia, distrito=excluded.distrito, direccion=excluded.direccion, fecha=excluded.fecha', [dni, 'Sí', *fields, ''])
        return page('Miembro de mesa guardado. Ya puedes descargar la lista en Excel.')
    except ValueError as exc:
        return page(str(exc), True, 400)

@app.get('/exportar')
def exportar():
    data = [[r['dni'], r['region'], r['provincia'], r['distrito'], r['direccion']] for r in rows() if r['estado']=='Sí']
    if not data:
        return page('Primero registra al menos un miembro de mesa para descargar el Excel.', True, 400)
    return excel(['DNI', 'Región', 'Provincia', 'Distrito', 'Dirección del local de votación'], data, 'miembros_de_mesa.xlsx')

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000)
